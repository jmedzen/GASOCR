import os
import zipfile
from pathlib import Path
from typing import List, Dict, Any
from docx import Document
from config import EXPORTS_DIR

def export_markdown(task: Dict[str, Any], pages: List[Dict[str, Any]]) -> Path:
    """匯出為 Markdown 格式"""
    task_id = task["id"]
    filename = Path(task["filename"]).stem
    output_path = EXPORTS_DIR / f"{filename}_{task_id[:8]}.md"
    
    lines = [
        f"# {filename} - OCR 轉譯成果",
        f"> 模型: `{task.get('model', 'gemini')}` | 語言: {task.get('lang_pref')} | 排版: {task.get('direction_pref')} / {task.get('column_pref')}",
        "",
    ]
    
    for page in pages:
        p_num = page["page_num"]
        ocr_content = page.get("ocr_text", "").strip()
        lines.append(f"\n---\n### 📄 第 {p_num} 頁\n")
        lines.append(ocr_content if ocr_content else "*(本頁未成功辨識或無文字)*")
        lines.append("")
        
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path

def export_txt(task: Dict[str, Any], pages: List[Dict[str, Any]]) -> Path:
    """匯出為純文字 TXT 格式"""
    task_id = task["id"]
    filename = Path(task["filename"]).stem
    output_path = EXPORTS_DIR / f"{filename}_{task_id[:8]}.txt"
    
    lines = [
        f"==================================================",
        f" 檔案名稱: {filename}",
        f" 轉譯模型: {task.get('model', 'gemini')}",
        f" 總頁數: {len(pages)}",
        f"==================================================\n",
    ]
    
    for page in pages:
        p_num = page["page_num"]
        ocr_content = page.get("ocr_text", "").strip()
        lines.append(f"\n-------------------- 第 {p_num} 頁 --------------------\n")
        lines.append(ocr_content if ocr_content else "(本頁無辨識文字)")
        lines.append("\n")
        
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path

def export_docx(task: Dict[str, Any], pages: List[Dict[str, Any]]) -> Path:
    """匯出為 Microsoft Word (.docx) 格式"""
    task_id = task["id"]
    filename = Path(task["filename"]).stem
    output_path = EXPORTS_DIR / f"{filename}_{task_id[:8]}.docx"
    
    doc = Document()
    doc.add_heading(f"{filename} - OCR 轉譯成果", level=0)
    
    # 加入元數據段落
    meta_p = doc.add_paragraph()
    meta_p.add_run(f"轉譯模型: {task.get('model')} | 語言設定: {task.get('lang_pref')}\n").italic = True
    
    for page in pages:
        p_num = page["page_num"]
        ocr_content = page.get("ocr_text", "").strip()
        
        doc.add_heading(f"第 {p_num} 頁", level=1)
        if ocr_content:
            for block in ocr_content.split("\n\n"):
                doc.add_paragraph(block.strip())
        else:
            doc.add_paragraph("(本頁無辨識文字)")
            
        doc.add_page_break()
        
    doc.save(str(output_path))
    return output_path

def export_zip_bundle(task: Dict[str, Any], pages: List[Dict[str, Any]]) -> Path:
    """打包匯出所有格式（MD, TXT, DOCX）以及頁面原圖壓縮包"""
    task_id = task["id"]
    filename = Path(task["filename"]).stem
    zip_path = EXPORTS_DIR / f"{filename}_{task_id[:8]}_full_bundle.zip"
    
    md_path = export_markdown(task, pages)
    txt_path = export_txt(task, pages)
    docx_path = export_docx(task, pages)
    
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(md_path, arcname=f"{filename}.md")
        zf.write(txt_path, arcname=f"{filename}.txt")
        zf.write(docx_path, arcname=f"{filename}.docx")
        
        # 將各頁原圖也納入 zip
        for page in pages:
            img_path = Path(page["image_path"])
            if img_path.exists():
                arcname = f"rendered_pages/page_{page['page_num']:04d}{img_path.suffix}"
                zf.write(img_path, arcname=arcname)
                
    return zip_path
